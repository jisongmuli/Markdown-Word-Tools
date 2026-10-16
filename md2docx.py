#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 Markdown 交流记录转换为排版良好的 DOCX。

用法: python md2docx.py <input.md> <output.docx>
"""

import re
import sys
import unicodedata

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"
CODE_FONT_EA = "微软雅黑"

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]*\]\([^)]*\)|<https?://[^>\s]+>)"
)


# ---------------------------------------------------------------- helpers
def shade(element, fill):
    """给 pPr / tcPr 等元素加底纹。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def set_run_fonts(run, latin, east_asian, size=None, bold=None,
                  italic=None, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asian)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def disp_width(text):
    """估算显示宽度：全角字符按 2 计。"""
    width = 0
    for ch in text:
        width += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return width


def code_font_size(lines):
    """根据最长行自适应字号，保证不折行。"""
    widest = max((disp_width(l) for l in lines), default=1)
    if widest <= 1:
        return 8.5
    available_pt = 17.0 * 28.35          # A4 宽 21cm - 左右各 2cm 页边距
    size = available_pt / (0.5525 * widest)
    return max(5.5, min(8.5, round(size, 1)))


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BBBBBB")
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


# ---------------------------------------------------------------- parsing
def is_block_start(line):
    s = line.rstrip()
    if s.strip() == "":
        return True
    if s.lstrip().startswith("```"):
        return True
    if re.match(r"^#{1,6}\s", s):
        return True
    if re.match(r"^\s*\|.*\|\s*$", s):
        return True
    if re.match(r"^\s*>\s?", s):
        return True
    if re.match(r"^\s*[-*+]\s+", s):
        return True
    if re.match(r"^\s*\d+\.\s+", s):
        return True
    if re.match(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$", s):
        return True
    return False


def split_table_row(row):
    """按未转义的 | 分割表格行。"""
    cells = re.split(r"(?<!\\)\|", row.strip())
    if cells and cells[0].strip() == "":
        cells = cells[1:]
    if cells and cells[-1].strip() == "":
        cells = cells[:-1]
    return [c.strip().replace("\\|", "|") for c in cells]


def parse_blocks(lines):
    blocks = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1                                    # 跳过结束围栏
            blocks.append(("code", "\n".join(buf).rstrip("\n")))
            continue

        if re.match(r"^\s*\|.*\|\s*$", line):
            rows = []
            while i < n and re.match(r"^\s*\|.*\|\s*$", lines[i]):
                rows.append(lines[i])
                i += 1
            blocks.append(("table", rows))
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line.strip())
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue

        if re.match(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$", line):
            blocks.append(("hr", None))
            i += 1
            continue

        if re.match(r"^\s*>\s?", line):
            buf = []
            while i < n and re.match(r"^\s*>\s?", lines[i]):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            blocks.append(("quote", buf))
            continue

        m = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
        if m:
            items = []
            while i < n:
                m2 = re.match(r"^(\s*)([-*+])\s+(.*)$", lines[i])
                if not m2:
                    break
                items.append((len(m2.group(1)) // 2, m2.group(3)))
                i += 1
            blocks.append(("ul", items))
            continue

        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            items = []
            while i < n:
                m2 = re.match(r"^(\s*)(\d+)\.\s+(.*)$", lines[i])
                if not m2:
                    break
                items.append((len(m2.group(1)) // 2, m2.group(3)))
                i += 1
            blocks.append(("ol", items))
            continue

        buf = []
        while i < n and not is_block_start(lines[i]):
            buf.append(lines[i].strip())
            i += 1
        if buf:
            blocks.append(("p", " ".join(buf)))
        else:
            i += 1

    return blocks


# ---------------------------------------------------------------- rendering
def add_inline(paragraph, text, base_size=10.5, base_color=None,
               force_italic=False, force_bold=False):
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            # 递归处理，支持 **加粗里嵌套 `代码`**
            add_inline(paragraph, token[2:-2], base_size, base_color,
                       force_italic, True)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            set_run_fonts(run, CODE_FONT, CODE_FONT_EA, base_size - 0.5,
                          bold=force_bold or None, italic=force_italic,
                          color="A31515" if base_color is None else base_color)
        elif token.startswith("[") and "](" in token:
            label = token[1:token.index("](")]
            run = paragraph.add_run(label)
            set_run_fonts(run, BODY_FONT, BODY_FONT, base_size,
                          bold=force_bold or None, italic=force_italic,
                          color=base_color)
        elif token.startswith("<") and token.endswith(">"):
            run = paragraph.add_run(token[1:-1])
            set_run_fonts(run, BODY_FONT, BODY_FONT, base_size,
                          bold=force_bold or None, italic=force_italic,
                          color="0563C1")
        else:
            run = paragraph.add_run(token)
            set_run_fonts(run, BODY_FONT, BODY_FONT, base_size,
                          bold=force_bold or None, italic=force_italic,
                          color=base_color)


def render_code(doc, code, indent_cm=0.0):
    lines = code.split("\n")
    size = code_font_size(lines)
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.left_indent = Cm(0.25 + indent_cm)
    pf.right_indent = Cm(0.1)
    pf.line_spacing = 1.0
    shade(p._p.get_or_add_pPr(), "F4F4F4")
    run = p.add_run()
    for idx, line in enumerate(lines):
        if idx:
            run.add_break()
        if line:
            run.add_text(line)
    set_run_fonts(run, CODE_FONT, CODE_FONT_EA, size)
    return p


def render_table(doc, rows):
    grid = []
    for row in rows:
        cells = split_table_row(row)
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue                                   # 分隔行
        grid.append(cells)
    if not grid:
        return
    ncols = max(len(r) for r in grid)
    grid = [r + [""] * (ncols - len(r)) for r in grid]

    table = doc.add_table(rows=len(grid), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for r, row in enumerate(grid):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.15
            add_inline(para, cell_text, base_size=9,
                       base_color="111111" if r else "000000")
            if r == 0:
                for run in para.runs:
                    run.font.bold = True
                shade(cell._tc.get_or_add_tcPr(), "E8EEF7")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_blocks(doc, blocks, indent_cm=0.0, in_quote=False, first=False):
    for kind, payload in blocks:
        if kind == "h":
            level, text = payload
            if first and level == 1:
                p = doc.add_paragraph(style="Title")
                add_inline(p, text, base_size=22)
                first = False
                continue
            p = doc.add_heading(level=min(level, 4))
            for run in list(p.runs):
                run._element.getparent().remove(run._element)
            add_inline(p, text, base_size={1: 17, 2: 14, 3: 12, 4: 11}[min(level, 4)])
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string("1F3864")
            first = False
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.35
            if indent_cm:
                p.paragraph_format.left_indent = Cm(indent_cm)
            add_inline(p, payload, base_size=10.5,
                       base_color="595959" if in_quote else None,
                       force_italic=in_quote)
            first = False
        elif kind == "ul":
            for depth, item in payload:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.left_indent = Cm(0.75 + 0.6 * depth + indent_cm)
                add_inline(p, item, base_size=10.5,
                           base_color="595959" if in_quote else None,
                           force_italic=in_quote)
            first = False
        elif kind == "ol":
            for depth, item in payload:
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.left_indent = Cm(0.75 + 0.6 * depth + indent_cm)
                add_inline(p, item, base_size=10.5,
                           base_color="595959" if in_quote else None,
                           force_italic=in_quote)
            first = False
        elif kind == "code":
            render_code(doc, payload, indent_cm)
            first = False
        elif kind == "table":
            render_table(doc, payload)
            first = False
        elif kind == "quote":
            inner = parse_blocks(payload)
            render_blocks(doc, inner, indent_cm=indent_cm + 0.5,
                          in_quote=True)
            first = False
        elif kind == "hr":
            add_horizontal_rule(doc)
            first = False


# ---------------------------------------------------------------- document
def build(md_path, docx_path):
    with open(md_path, encoding="utf-8") as fh:
        text = fh.read()

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # 页脚页码
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    set_run_fonts(run, BODY_FONT, BODY_FONT, 9, color="808080")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._element.addnext(fld)

    blocks = parse_blocks(text.split("\n"))
    render_blocks(doc, blocks, first=True)

    doc.save(docx_path)
    print("saved:", docx_path)
    print("blocks:", len(blocks))


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Convert UTF-8 Markdown to a styled Word document.")
    parser.add_argument("input", help="Input Markdown file")
    parser.add_argument("output", help="Output .docx file")
    args = parser.parse_args()
    build(args.input, args.output)
